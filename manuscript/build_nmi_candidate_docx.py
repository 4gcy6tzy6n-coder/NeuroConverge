from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'manuscript' / 'ARTICLE_DRAFT.md'
OUTPUT = ROOT / 'manuscript' / 'NMI_MANUSCRIPT_CANDIDATE.docx'
FIGURES = ROOT / 'figures'

def set_font(run, name='Cambria', size=10.5, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)

def add_formatted(paragraph, text, size=10.5):
    # Lightweight Markdown inline formatting, including links and math delimiters.
    token = re.compile(r'(\*\*.+?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\$[^$]+\$)')
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=size)
        value = match.group(0)
        if value.startswith('**'):
            set_font(paragraph.add_run(value[2:-2]), size=size, bold=True)
        elif value.startswith('*'):
            set_font(paragraph.add_run(value[1:-1]), size=size, italic=True)
        elif value.startswith('`'):
            set_font(paragraph.add_run(value[1:-1]), name='Cambria Math', size=size)
        elif value.startswith('['):
            link = re.match(r'\[([^\]]+)\]\(([^)]+)\)', value)
            set_font(paragraph.add_run(link.group(1)), size=size, color=(37, 86, 132))
        else:
            set_font(paragraph.add_run(equation_text(value[1:-1])), name='Cambria Math', size=size)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=size)

def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar'); fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar'); fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1); run._r.append(instr); run._r.append(fld_char2)
    set_font(run, size=9, color=(105, 115, 125))

def equation_text(raw):
    raw = raw.strip().replace('\\mathrm{out\\_strength}_i', 'out-strengthᵢ')
    raw = raw.replace('\\mathrm{in\\_strength}_i', 'in-strengthᵢ')
    raw = raw.replace('\\sum_j W_{ij}W_{ji}', 'Σⱼ WᵢⱼWⱼᵢ')
    raw = raw.replace('\\frac{Σⱼ WᵢⱼWⱼᵢ}{(out-strengthᵢ)(in-strengthᵢ)}', 'Σⱼ WᵢⱼWⱼᵢ / [(out-strengthᵢ)(in-strengthᵢ)]')
    raw = raw.replace('h_t', 'hₜ').replace('h_{t-1}', 'hₜ₋₁').replace('W_{hh}', 'Wₕₕ').replace('W_x', 'Wₓ').replace('x_t', 'xₜ')
    raw = raw.replace('L_i', 'Lᵢ').replace('W_{ij}', 'Wᵢⱼ').replace('W_{ji}', 'Wⱼᵢ')
    raw = raw.replace('\\lambda', 'λ').replace('\\tanh', 'tanh').replace('\\sum_j', 'Σⱼ')
    raw = raw.replace('\\mathrm{', '').replace('}', '').replace('{', '').replace('\\mathsf{T}', 'ᵀ')
    raw = raw.replace('\\left', '').replace('\\right', '').replace('\\cdot', '·')
    raw = raw.replace('\\times', '×').replace('\\epsilon', 'ε').replace('\\in', '∈')
    raw = raw.replace('\\', '')
    return raw

def add_paragraph(doc, text, style=None, size=10.5, before=0, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_together = keep
    add_formatted(p, text, size=size)
    return p

md = SOURCE.read_text(encoding='utf-8').splitlines()
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27); section.page_height = Inches(11.69)
section.top_margin = Inches(.72); section.bottom_margin = Inches(.72)
section.left_margin = Inches(.78); section.right_margin = Inches(.78)
styles = doc.styles
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Cambria'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Cambria')
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Cambria')
# The built-in Word Title style carries a paragraph rule in some templates.
# Remove it so the title renders as plain text in all Word-compatible viewers.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)
styles['Normal'].font.size = Pt(10.5)
styles['Heading 1'].font.size = Pt(13); styles['Heading 1'].font.bold = True
styles['Heading 1'].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].paragraph_format.space_before = Pt(13)
styles['Heading 1'].paragraph_format.space_after = Pt(5)
styles['Heading 2'].font.size = Pt(11); styles['Heading 2'].font.bold = True
styles['Heading 2'].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(4)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run('NeuroConverge | manuscript candidate'), size=8, color=(0, 0, 0))
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('Internal candidate  •  '), size=9, color=(105, 115, 125)); add_page_field(footer)

fig_legend_lines = []
inside_legend = False
in_equation = False
math_lines = []
seen_title = False
for line in md:
    stripped = line.strip()
    if not stripped:
        continue
    if stripped == '## Figure legends':
        inside_legend = True
        p = doc.add_paragraph('Figure legends', style='Heading 1')
        continue
    if inside_legend and stripped.startswith('**Figure '):
        fig_legend_lines.append(stripped)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_together = True
        add_formatted(p, stripped, size=9.5)
        continue
    if stripped.startswith('# '):
        if not seen_title:
            title = stripped[2:]
            p = doc.add_paragraph(style='Title')
            p.paragraph_format.space_after = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(title), size=17, bold=True, color=(0, 0, 0))
            seen_title = True
        continue
    if stripped.startswith('**NeuroConverge internal manuscript candidate'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        add_formatted(p, stripped.replace('**', ''), size=9)
        continue
    if stripped == '## Abstract':
        doc.add_paragraph('Abstract', style='Heading 1')
        continue
    if stripped.startswith('## '):
        heading = stripped[3:]
        if heading == 'References':
            doc.add_page_break()
        doc.add_paragraph(heading, style='Heading 1')
        continue
    if stripped.startswith('### '):
        doc.add_paragraph(stripped[4:], style='Heading 2')
        continue
    if stripped == r'\[':
        in_equation = True; math_lines = []; continue
    if stripped == r'\]':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
        set_font(p.add_run(equation_text(' '.join(math_lines))), name='Cambria Math', size=10)
        in_equation = False; continue
    if in_equation:
        math_lines.append(stripped); continue
    if re.match(r'^\d+\.\s', stripped):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.24)
        p.paragraph_format.first_line_indent = Inches(-.24)
        p.paragraph_format.space_after = Pt(5)
        add_formatted(p, stripped, size=9.5)
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    add_formatted(p, stripped, size=10.5)

# Add publication-sized figures at the end, one per page, with the same captions as the manuscript.
figure_map = [
    ('Fig1_motivation', 'Fig1_motivation.png'),
    ('Fig2_framework', 'Fig2_framework.png'),
    ('Fig3_structure_boundary', 'Fig3_structure_boundary.png'),
    ('Fig4_motor_feedback', 'Fig4_motor_feedback.png'),
    ('Fig5_temporal_credit', 'Fig5_temporal_credit.png'),
    ('Fig6_evidence_map', 'Fig6_evidence_map.png'),
]
if len(fig_legend_lines) != len(figure_map):
    raise ValueError(f'Expected six figure legends, found {len(fig_legend_lines)}')
doc.add_page_break()
doc.add_paragraph('Figures', style='Heading 1')
for index, ((directory, filename), caption) in enumerate(zip(figure_map, fig_legend_lines)):
    if index:
        doc.add_page_break()
    image_path = FIGURES / directory / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(image_path), width=Inches(6.65))
    cap = doc.add_paragraph(); cap.paragraph_format.space_after = Pt(0)
    add_formatted(cap, caption, size=8.5)

doc.core_properties.title = 'Evidence boundaries in transferring biological computation to artificial systems'
doc.core_properties.subject = 'Nature Machine Intelligence-style internal manuscript candidate'
doc.core_properties.author = 'NeuroConverge project authorship pending'
doc.save(OUTPUT)
print(OUTPUT)
